"""
SERP Analyzer para Linkbuilding - Versión Playwright v3
España, clasificación automática de intención y soporte AI Overview.
"""

from flask import Flask, render_template, request, jsonify, send_file
import time
import random
import re
import io
import os
from urllib.parse import urlparse, unquote, quote_plus
from datetime import datetime
from collections import defaultdict
import json
import html as html_mod
from bs4 import BeautifulSoup, Comment
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from playwright.sync_api import sync_playwright

app = Flask(__name__)

# Fichero donde se guardan cookies/localStorage entre sesiones
BROWSER_STATE_FILE = "browser_state.json"

MOBILE_CONFIG = {
    "viewport": {"width": 412, "height": 915},
    "user_agent": (
        "Mozilla/5.0 (Linux; Android 14; SM-S928B) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/131.0.6778.200 Mobile Safari/537.36"
    ),
    "device_scale_factor": 2.625,
    "is_mobile": True,
    "has_touch": True,
}

# Script inyectado antes de cada página para ocultar señales de automatización
_STEALTH_SCRIPT = """
    Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
    Object.defineProperty(navigator, 'plugins', { get: () => [1, 2, 3, 4, 5] });
    Object.defineProperty(navigator, 'languages', { get: () => ['es-ES', 'es'] });
    window.chrome = { runtime: {} };
"""

COUNTRY = {
    "name": "España",
    "domain": "google.es",
    "gl": "es",
    "hl": "es",
    "geo": {"latitude": 40.4168, "longitude": -3.7038},
    "timezone": "Europe/Madrid",
}


def extract_domain(url):
    try:
        parsed = urlparse(url)
        domain = parsed.netloc
        if domain.startswith("www."):
            domain = domain[4:]
        return domain
    except:
        return None


def clean_url(url):
    if not url:
        return None
    url = unquote(url)
    if "/url?q=" in url:
        match = re.search(r'/url\?q=([^&]+)', url)
        if match:
            url = unquote(match.group(1))
    if url.startswith("/search") or url.startswith("/"):
        return None
    if "google." in url and "/search" in url:
        return None
    return url


def detect_serp_features(soup):
    """Detecta las features presentes en la SERP."""
    features = {
        "google_ads": False,
        "ai_overview": False,
        "featured_snippet": False,
        "people_also_ask": False,
        "image_carousel": False,
        "video_carousel": False,
        "local_pack": False,
        "shopping_results": False,
        "knowledge_panel": False,
        "news": False,
    }

    html_str = str(soup).lower()

    if soup.find("span", string=re.compile(r"^(Patrocinado|Sponsored|Anuncio)$", re.I)) or soup.find("div", {"data-text-ad": True}):
        features["google_ads"] = True

    if soup.find("div", {"data-async-context-required": "aobs"}) or "ai overview" in html_str or "generado por ia" in html_str or "descripción general de ia" in html_str:
        features["ai_overview"] = True

    if soup.find("div", {"class": "xpdopen"}):
        features["featured_snippet"] = True

    if soup.find("div", {"jsname": "Cpkphb"}) or "preguntas relacionadas" in html_str or "people also ask" in html_str:
        features["people_also_ask"] = True

    if soup.find("div", {"id": "iur"}) or soup.find("g-scrolling-carousel"):
        features["image_carousel"] = True

    if soup.find("video-voyager") or "youtube.com/watch" in html_str:
        features["video_carousel"] = True

    if soup.find("div", {"class": re.compile(r"VkpGBb")}) or soup.find("a", {"href": re.compile(r"maps\.google")}):
        features["local_pack"] = True

    if soup.find("div", {"class": re.compile(r"commercial-unit|cu-container")}) or "google.es/shopping" in html_str or "google.com/shopping" in html_str:
        features["shopping_results"] = True

    if soup.find("div", {"class": re.compile(r"kp-wholepage|liYKde")}):
        features["knowledge_panel"] = True

    if soup.find("div", {"class": re.compile(r"ftSUBd|So9e7d")}) or "noticias principales" in html_str or "top stories" in html_str:
        features["news"] = True

    return features


def classify_intent(features):
    """Clasifica la intención de búsqueda según las features SERP detectadas.

    Pesos:
      Transaccional: shopping (+2), ads (+1), local_pack (+1)
      Informacional: featured_snippet (+2), paa (+1), news (+1), ai_overview (+1), video (+1)
      Branding:      knowledge_panel (+3)
    """
    t = 0
    i = 0
    b = 0

    if features.get("shopping_results"):
        t += 2
    if features.get("google_ads"):
        t += 1
    if features.get("local_pack"):
        t += 1

    if features.get("featured_snippet"):
        i += 2
    if features.get("people_also_ask"):
        i += 1
    if features.get("news"):
        i += 1
    if features.get("ai_overview"):
        i += 1
    if features.get("video_carousel"):
        i += 1

    if features.get("knowledge_panel"):
        b += 3

    if b >= 3 and t == 0 and i <= 1:
        return "branding"
    if t > 0 and i > 0:
        return "mixto"
    if t > i:
        return "transaccional"
    if i > t:
        return "informacional"
    return "mixto"


def extract_organic_results(soup):
    """Extrae los resultados orgánicos del HTML."""
    organic_results = []
    seen_urls = set()
    position = 0

    print("[DEBUG] Iniciando extracción de resultados orgánicos...")

    result_links = soup.find_all("a", class_="rTyHce")
    print(f"[DEBUG] Enlaces rTyHce encontrados: {len(result_links)}")

    for link in result_links:
        if position >= 10:
            break

        href = link.get("href", "")
        if not href:
            continue

        url = clean_url(href)
        if not url or url in seen_urls:
            continue

        domain = extract_domain(url)
        if not domain or "google." in domain:
            continue

        is_ad = False
        parent = link
        for _ in range(8):
            parent = parent.find_parent()
            if not parent:
                break
            parent_text = parent.get_text()
            if any(ad_text in parent_text for ad_text in ["Patrocinado", "Sponsored", "Anuncio", "Resultado patrocinado"]):
                is_ad = True
                break
            if parent.get("data-text-ad") or parent.get("data-ae"):
                is_ad = True
                break

        if is_ad:
            continue

        title = None
        heading_div = link.find("div", attrs={"role": "heading"})
        if heading_div:
            title_span = heading_div.find("span")
            title = title_span.get_text(strip=True) if title_span else heading_div.get_text(strip=True)

        if not title:
            title_div = link.find("div", class_=re.compile(r"F0FGWb|v7jaNc|MBeuO"))
            if title_div:
                title = title_div.get_text(strip=True)

        if not title:
            h3 = link.find("h3")
            if h3:
                title = h3.get_text(strip=True)

        if not title or len(title) < 5:
            continue

        seen_urls.add(url)
        position += 1
        organic_results.append({"position": position, "title": title, "url": url, "domain": domain})
        print(f"[DEBUG] Resultado {position}: {domain} - {title[:50]}...")

    if len(organic_results) == 0:
        print("[DEBUG] Intentando estrategia alternativa con MjjYud...")
        containers = soup.find_all("div", class_="MjjYud")
        print(f"[DEBUG] Contenedores MjjYud encontrados: {len(containers)}")

        for container in containers:
            if position >= 10:
                break

            link = container.find("a", href=True)
            if not link:
                continue

            href = link.get("href", "")
            url = clean_url(href)
            if not url or url in seen_urls:
                continue

            domain = extract_domain(url)
            if not domain or "google." in domain:
                continue

            title = None
            heading = container.find(attrs={"role": "heading"})
            if heading:
                title = heading.get_text(strip=True)
            if not title:
                h3 = container.find("h3")
                if h3:
                    title = h3.get_text(strip=True)

            if not title or len(title) < 5:
                continue

            container_text = container.get_text()
            if any(ad in container_text for ad in ["Patrocinado", "Sponsored", "Anuncio"]):
                continue

            seen_urls.add(url)
            position += 1
            organic_results.append({"position": position, "title": title, "url": url, "domain": domain})
            print(f"[DEBUG] Resultado {position}: {domain} - {title[:50]}...")

    print(f"[DEBUG] Total resultados orgánicos: {len(organic_results)}")
    return organic_results


def _handle_cookie_consent(page):
    """Acepta el consentimiento de cookies de Google si aparece."""
    try:
        page.wait_for_timeout(1500)
        for selector in [
            'button:text("Aceptar todo")',
            'button:text("Accept all")',
            '[aria-label="Aceptar todo"]',
            'button[jsname="b3VHJd"]',
        ]:
            try:
                btn = page.locator(selector).first
                if btn.is_visible(timeout=1500):
                    btn.click()
                    page.wait_for_timeout(1200)
                    print("[INFO] Consentimiento de cookies aceptado")
                    return
            except Exception:
                continue
    except Exception:
        pass


def _decode_tgqphd_url(url):
    """Decodifica escapes \\uXXXX residuales en URLs (doble-escape que sobrevive al json.loads)."""
    return re.sub(r"\\u([0-9a-fA-F]{4})", lambda m: chr(int(m.group(1), 16)), url)


def _url_has_path(url):
    """Devuelve True si la URL tiene un path significativo (no solo dominio)."""
    try:
        from urllib.parse import urlparse
        path = urlparse(url).path
        return bool(path and path not in ("", "/"))
    except Exception:
        return False


def extract_ai_overview_results(html_content):
    """Extrae fuentes del AI Overview usando regex sobre el HTML crudo.

    Google embebe los datos de citas en dos formatos dentro de marcadores TgQPHd:
      Formato A (cabecera, len≥12): [null, null, uuid, ..., site_name(8), ..., url(11), ...]
      Formato C (cita expandida): [uuid_str, pos_num_str, 0|1, url, ...]  len=7

    Los marcadores aparecen tanto como comentarios HTML reales (<!--TgQPHd|...-->)
    como embebidos en strings JS serializados (\x3c!--TgQPHd|...-->). BeautifulSoup
    solo encuentra los primeros; usamos regex sobre el HTML bruto para capturar todos.
    """
    ai_results = []
    seen_urls = set()
    seen_domains = set()
    uuid_to_name = {}  # UUID → site_name (solo para registrar; no se propaga a otras citas)

    raw_matches = re.findall(r"TgQPHd\|(\[.*?\])(?=-->)", html_content, re.DOTALL)

    parsed_items = []
    for raw in raw_matches:
        if raw == "[]":
            continue
        try:
            data = json.loads(html_mod.unescape(raw))
        except (json.JSONDecodeError, ValueError):
            continue
        if not isinstance(data, list) or not data:
            continue
        item = data[0] if isinstance(data[0], list) else data
        if isinstance(item, list):
            parsed_items.append(item)

    def _add(url, site_name):
        """Añade al resultado si la URL es válida y no está duplicada."""
        if len(ai_results) >= 25:
            return False
        url = _decode_tgqphd_url(url)
        if not url.startswith("http"):
            return False
        domain = extract_domain(url)
        if not domain or "google." in domain:
            return False
        # Si es solo dominio (sin path), descartar si ya tenemos esa fuente
        if not _url_has_path(url) and domain in seen_domains:
            return False
        if url in seen_urls:
            return False
        seen_urls.add(url)
        seen_domains.add(domain)
        ai_results.append({
            "position": len(ai_results) + 1,
            "title": site_name,
            "url": url,
            "domain": domain,
        })
        return True

    # — Paso 1: Formato A — cabecera principal con site_name —
    for item in parsed_items:
        if len(ai_results) >= 25:
            break
        if (
            len(item) >= 12
            and item[0] is None
            and isinstance(item[8], str) and item[8]
            and isinstance(item[11], str) and item[11].startswith("http")
        ):
            uuid = item[2] if len(item) > 2 and isinstance(item[2], str) else None
            if uuid:
                uuid_to_name[uuid] = item[8]
            _add(item[11], item[8])

    # — Paso 2: Formato C — citas expandidas [uuid, pos_str, flag, url, ...] len=7 —
    for item in parsed_items:
        if len(ai_results) >= 25:
            break
        # Formato C: [uuid_str, pos_num_str, flag_int, url, ...]
        # len entre 4 y 10, item[1] es string numérico (posición de la cita)
        if (
            4 <= len(item) <= 10
            and isinstance(item[0], str)
            and isinstance(item[1], str) and item[1].isdigit()
            and isinstance(item[3], str) and item[3].startswith("http")
        ):
            domain = extract_domain(_decode_tgqphd_url(item[3]))
            _add(item[3], domain or item[0][:30])

    print(f"[DEBUG] Fuentes AI Overview (TgQPHd): {len(ai_results)}")
    return ai_results


def _expand_ai_citations(page):
    """Pulsa los botones de citas expandibles del AI Overview para cargar todas las fuentes.

    Google carga solo la primera fuente de cada grupo en el HTML inicial.
    Los botones 'Wikipedia (+3)' etc. tienen un atributo disabled que bloquea la propagación
    del evento click — el handler real está en el span padre con jsaction. La solución es
    quitar disabled con JS y disparar el click sobre el span padre directamente.
    """
    try:
        # Esperar a que jsaction inicialice los botones
        try:
            page.wait_for_function(
                "() => document.querySelectorAll('button[data-amic=\"true\"]').length > 0",
                timeout=3000,
            )
        except Exception:
            pass

        clicked = page.evaluate("""
            () => {
                const buttons = Array.from(document.querySelectorAll('button[data-amic="true"]'));
                let count = 0;
                buttons.forEach(btn => {
                    try {
                        btn.removeAttribute('disabled');
                        // El handler jsaction está en el span padre más cercano con jsaction
                        const target = btn.closest('span[jsaction]') || btn.parentElement || btn;
                        target.dispatchEvent(new MouseEvent('click', {
                            bubbles: true,
                            cancelable: true,
                            view: window
                        }));
                        count++;
                    } catch(e) {}
                });
                return count;
            }
        """)

        if clicked and clicked > 0:
            print(f"[DEBUG] Expandiendo {clicked} grupos de citas AI Overview...")
            # Esperar proporcional al número de grupos para que se carguen las fuentes
            page.wait_for_timeout(min(clicked * 2000, 8000))

        return clicked or 0

    except Exception as e:
        print(f"[DEBUG] Error al expandir citas AI: {e}")
        return 0


def _scrape_keyword_on_page(page, keyword):
    """Extrae todos los datos de una keyword sobre una página Playwright ya abierta."""
    result = {
        "keyword": keyword,
        "country": COUNTRY["name"],
        "timestamp": datetime.now().isoformat(),
        "organic_results": [],
        "ai_overview_results": [],
        "serp_features": {},
        "error": None,
    }

    try:
        encoded_keyword = quote_plus(keyword)
        search_url = (
            f"https://www.{COUNTRY['domain']}/search"
            f"?q={encoded_keyword}&hl={COUNTRY['hl']}&gl={COUNTRY['gl']}"
        )

        page.goto(search_url, wait_until="networkidle")
        # Espera extra para que AI Overview cargue (se inyecta via JS tras el HTML inicial)
        page.wait_for_timeout(3000)

        # Scroll suave para activar lazy content
        page.evaluate("window.scrollBy(0, 600)")
        page.wait_for_timeout(800)
        page.evaluate("window.scrollBy(0, -200)")
        page.wait_for_timeout(500)

        # Expandir citas del AI Overview antes de leer el HTML
        n_expanded = _expand_ai_citations(page)
        if n_expanded:
            # Esperar a que los comentarios TgQPHd adicionales aparezcan en el DOM
            try:
                page.wait_for_function(
                    """() => {
                        const html = document.documentElement.innerHTML;
                        const matches = html.match(/TgQPHd\|/g);
                        return matches && matches.length > 1;
                    }""",
                    timeout=5000,
                )
            except Exception:
                page.wait_for_timeout(2000)

        html_content = page.content()

        safe_kw = re.sub(r"[^\w]", "_", keyword)[:30]
        with open(f"debug_{safe_kw}.html", "w", encoding="utf-8") as f:
            f.write(html_content)

        if "captcha" in html_content.lower() or "unusual traffic" in html_content.lower():
            result["error"] = "CAPTCHA detectado."
            return result

        soup = BeautifulSoup(html_content, "html.parser")
        result["serp_features"] = detect_serp_features(soup)
        result["organic_results"] = extract_organic_results(soup)
        # Las fuentes de AI Overview se extraen con regex del HTML crudo (no de BS4)
        result["ai_overview_results"] = extract_ai_overview_results(html_content)

        print(
            f"[DEBUG] '{keyword}': orgánicos={len(result['organic_results'])}, "
            f"AI fuentes={len(result['ai_overview_results'])}"
        )

        if not result["organic_results"]:
            result["error"] = "No se encontraron resultados orgánicos."

    except Exception as e:
        result["error"] = f"Error: {str(e)}"
        import traceback
        traceback.print_exc()

    return result


def scrape_all_keywords(keywords_list, delay_min=4, delay_max=7):
    """Lanza una única sesión de navegador para todas las keywords."""
    all_results = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=False,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--disable-dev-shm-usage",
                "--no-sandbox",
                "--disable-infobars",
                "--disable-notifications",
            ],
        )

        # Cargar estado guardado (cookies, localStorage) si existe
        ctx_kwargs = dict(
            viewport=MOBILE_CONFIG["viewport"],
            user_agent=MOBILE_CONFIG["user_agent"],
            device_scale_factor=MOBILE_CONFIG["device_scale_factor"],
            is_mobile=MOBILE_CONFIG["is_mobile"],
            has_touch=MOBILE_CONFIG["has_touch"],
            geolocation=COUNTRY["geo"],
            permissions=["geolocation"],
            locale=f"{COUNTRY['hl']}-{COUNTRY['gl'].upper()}",
            timezone_id=COUNTRY["timezone"],
        )
        if os.path.exists(BROWSER_STATE_FILE):
            ctx_kwargs["storage_state"] = BROWSER_STATE_FILE
            print("[INFO] Estado de sesión anterior cargado")

        context = browser.new_context(**ctx_kwargs)

        # Inyectar parches anti-detección antes de cada navegación
        context.add_init_script(_STEALTH_SCRIPT)

        page = context.new_page()
        page.set_default_timeout(30000)

        # Warm-up: visitar Google home, aceptar cookies y guardar estado
        print("[INFO] Iniciando sesión en Google España...")
        page.goto(f"https://www.{COUNTRY['domain']}", wait_until="networkidle")
        _handle_cookie_consent(page)
        page.wait_for_timeout(random.randint(2000, 3500))

        # Persistir cookies/localStorage para la próxima ejecución
        context.storage_state(path=BROWSER_STATE_FILE)
        print("[INFO] Estado de sesión guardado")

        for i, kw_data in enumerate(keywords_list):
            keyword = kw_data["keyword"]
            print(f"\n[INFO] Keyword {i + 1}/{len(keywords_list)}: '{keyword}'")

            if i > 0:
                delay = random.uniform(delay_min, delay_max)
                print(f"[INFO] Esperando {delay:.1f}s...")
                time.sleep(delay)

            result = _scrape_keyword_on_page(page, keyword)
            all_results.append(result)

            # Actualizar estado guardado tras cada keyword exitosa
            if not result.get("error"):
                context.storage_state(path=BROWSER_STATE_FILE)

        browser.close()

    return all_results


def aggregate_domains(all_results):
    """Agrega dominios de resultados orgánicos."""
    domain_stats = defaultdict(lambda: {"appearances": 0, "keywords": [], "positions": [], "urls": set()})

    for result in all_results:
        if result.get("error") and not result.get("organic_results"):
            continue
        keyword = result["keyword"]
        for organic in result.get("organic_results", []):
            domain = organic["domain"]
            domain_stats[domain]["appearances"] += 1
            domain_stats[domain]["keywords"].append(keyword)
            domain_stats[domain]["positions"].append(organic["position"])
            domain_stats[domain]["urls"].add(organic["url"])

    aggregated = []
    for domain, stats in domain_stats.items():
        avg_position = sum(stats["positions"]) / len(stats["positions"]) if stats["positions"] else 0
        aggregated.append({
            "domain": domain,
            "appearances": stats["appearances"],
            "avg_position": round(avg_position, 2),
            "keywords": list(set(stats["keywords"])),
            "unique_urls": len(stats["urls"]),
            "urls": list(stats["urls"]),
        })

    aggregated.sort(key=lambda x: (-x["appearances"], x["avg_position"]))
    return aggregated


def aggregate_ai_domains(all_results):
    """Agrega dominios de resultados de AI Overview."""
    domain_stats = defaultdict(lambda: {"appearances": 0, "keywords": [], "positions": [], "urls": set()})

    for result in all_results:
        if not result.get("ai_overview_results"):
            continue
        keyword = result["keyword"]
        for item in result.get("ai_overview_results", []):
            domain = item["domain"]
            domain_stats[domain]["appearances"] += 1
            domain_stats[domain]["keywords"].append(keyword)
            domain_stats[domain]["positions"].append(item["position"])
            domain_stats[domain]["urls"].add(item["url"])

    aggregated = []
    for domain, stats in domain_stats.items():
        avg_position = sum(stats["positions"]) / len(stats["positions"]) if stats["positions"] else 0
        aggregated.append({
            "domain": domain,
            "appearances": stats["appearances"],
            "avg_position": round(avg_position, 2),
            "keywords": list(set(stats["keywords"])),
            "unique_urls": len(stats["urls"]),
            "urls": list(stats["urls"]),
        })

    aggregated.sort(key=lambda x: (-x["appearances"], x["avg_position"]))
    return aggregated


def parse_keywords_input(text):
    """Parsea el input de keywords (una por línea)."""
    keywords = []
    for line in text.strip().split("\n"):
        keyword = line.strip()
        if keyword:
            keywords.append({"keyword": keyword})
    return keywords


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/analyze", methods=["POST"])
def analyze():
    data = request.get_json()

    if not data or "keywords" not in data:
        return jsonify({"error": "No se proporcionaron keywords"}), 400

    keywords_input = data.get("keywords", "")
    delay_min = data.get("delay_min", 4)
    delay_max = data.get("delay_max", 7)

    keywords_list = parse_keywords_input(keywords_input)

    if not keywords_list:
        return jsonify({"error": "No se encontraron keywords válidas"}), 400

    all_results = scrape_all_keywords(keywords_list, delay_min=delay_min, delay_max=delay_max)

    for result in all_results:
        result["tag"] = classify_intent(result.get("serp_features", {}))

    domain_aggregation = aggregate_domains(all_results)
    ai_domain_aggregation = aggregate_ai_domains(all_results)

    return jsonify({
        "success": True,
        "results": all_results,
        "domain_aggregation": domain_aggregation,
        "ai_domain_aggregation": ai_domain_aggregation,
        "total_keywords": len(keywords_list),
        "keywords_with_errors": sum(1 for r in all_results if r.get("error")),
    })


@app.route("/api/export/xlsx", methods=["POST"])
def export_xlsx():
    data = request.get_json()

    if not data:
        return jsonify({"error": "No hay datos para exportar"}), 400

    results = data.get("results", [])
    domain_aggregation = data.get("domain_aggregation", [])
    ai_domain_aggregation = data.get("ai_domain_aggregation", [])

    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill_green = PatternFill("solid", fgColor="2E7D32")
    header_fill_blue = PatternFill("solid", fgColor="1565C0")
    alt_fill = PatternFill("solid", fgColor="F5F5F5")
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    def apply_header(ws, headers, fill):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

    def apply_row_format(ws, row_num, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            if row_num % 2 == 0:
                cell.fill = alt_fill

    # Hoja 1: Resultados orgánicos por keyword
    ws1 = wb.active
    ws1.title = "Resultados por Keyword"

    headers1 = ["Keyword", "Intención", "País", "Posición", "Título", "URL", "Dominio",
                "Ads", "AI Overview", "Featured Snippet", "PAA", "Imágenes",
                "Vídeos", "Local Pack", "Shopping", "Knowledge Panel", "Noticias", "Error"]
    apply_header(ws1, headers1, header_fill_green)

    row_num = 2
    for result in results:
        keyword = result.get("keyword", "")
        tag = result.get("tag", "")
        country = result.get("country", "")
        error = result.get("error", "")
        features = result.get("serp_features", {})
        organic = result.get("organic_results", [])

        feat_vals = [
            "Sí" if features.get("google_ads") else "No",
            "Sí" if features.get("ai_overview") else "No",
            "Sí" if features.get("featured_snippet") else "No",
            "Sí" if features.get("people_also_ask") else "No",
            "Sí" if features.get("image_carousel") else "No",
            "Sí" if features.get("video_carousel") else "No",
            "Sí" if features.get("local_pack") else "No",
            "Sí" if features.get("shopping_results") else "No",
            "Sí" if features.get("knowledge_panel") else "No",
            "Sí" if features.get("news") else "No",
        ]

        rows_to_write = organic if organic else [None]
        for i, org in enumerate(rows_to_write):
            ws1.cell(row=row_num, column=1, value=keyword if i == 0 else "")
            ws1.cell(row=row_num, column=2, value=tag if i == 0 else "")
            ws1.cell(row=row_num, column=3, value=country if i == 0 else "")
            if org:
                ws1.cell(row=row_num, column=4, value=org.get("position", ""))
                ws1.cell(row=row_num, column=5, value=org.get("title", ""))
                ws1.cell(row=row_num, column=6, value=org.get("url", ""))
                ws1.cell(row=row_num, column=7, value=org.get("domain", ""))
            if i == 0:
                for j, val in enumerate(feat_vals, 8):
                    ws1.cell(row=row_num, column=j, value=val)
                ws1.cell(row=row_num, column=18, value=error)
            apply_row_format(ws1, row_num, 18)
            row_num += 1

    for i, width in enumerate([30, 15, 12, 8, 50, 60, 25, 6, 10, 14, 6, 8, 8, 10, 8, 14, 8, 30], 1):
        ws1.column_dimensions[get_column_letter(i)].width = width

    # Hoja 2: Dominios orgánicos agregados
    ws2 = wb.create_sheet("Dominios Orgánicos")
    headers2 = ["#", "Dominio", "Apariciones", "Posición Media", "URLs Únicas", "Keywords"]
    apply_header(ws2, headers2, header_fill_green)

    for idx, dom in enumerate(domain_aggregation, 1):
        row = idx + 1
        ws2.cell(row=row, column=1, value=idx)
        ws2.cell(row=row, column=2, value=dom.get("domain", ""))
        ws2.cell(row=row, column=3, value=dom.get("appearances", 0))
        ws2.cell(row=row, column=4, value=dom.get("avg_position", 0))
        ws2.cell(row=row, column=5, value=dom.get("unique_urls", 0))
        ws2.cell(row=row, column=6, value=", ".join(dom.get("keywords", [])))
        apply_row_format(ws2, row, 6)

    for i, width in enumerate([5, 35, 12, 14, 12, 80], 1):
        ws2.column_dimensions[get_column_letter(i)].width = width

    # Hoja 3: Resultados AI Overview por keyword
    ws3 = wb.create_sheet("AI Overview por Keyword")
    headers3 = ["Keyword", "Intención", "Posición IA", "Título", "URL", "Dominio", "Error"]
    apply_header(ws3, headers3, header_fill_blue)

    row_num3 = 2
    for result in results:
        keyword = result.get("keyword", "")
        tag = result.get("tag", "")
        error = result.get("error", "")
        ai_res = result.get("ai_overview_results", [])

        rows_to_write = ai_res if ai_res else [None]
        for i, ai in enumerate(rows_to_write):
            ws3.cell(row=row_num3, column=1, value=keyword if i == 0 else "")
            ws3.cell(row=row_num3, column=2, value=tag if i == 0 else "")
            if ai:
                ws3.cell(row=row_num3, column=3, value=ai.get("position", ""))
                ws3.cell(row=row_num3, column=4, value=ai.get("title", ""))
                ws3.cell(row=row_num3, column=5, value=ai.get("url", ""))
                ws3.cell(row=row_num3, column=6, value=ai.get("domain", ""))
            if i == 0:
                ws3.cell(row=row_num3, column=7, value=error if error else ("Sin AI Overview" if not ai_res else ""))
            apply_row_format(ws3, row_num3, 7)
            row_num3 += 1

    for i, width in enumerate([30, 15, 12, 50, 60, 25, 30], 1):
        ws3.column_dimensions[get_column_letter(i)].width = width

    # Hoja 4: Dominios AI Overview agregados
    ws4 = wb.create_sheet("Dominios AI Overview")
    headers4 = ["#", "Dominio", "Apariciones", "Posición Media IA", "URLs Únicas", "Keywords"]
    apply_header(ws4, headers4, header_fill_blue)

    for idx, dom in enumerate(ai_domain_aggregation, 1):
        row = idx + 1
        ws4.cell(row=row, column=1, value=idx)
        ws4.cell(row=row, column=2, value=dom.get("domain", ""))
        ws4.cell(row=row, column=3, value=dom.get("appearances", 0))
        ws4.cell(row=row, column=4, value=dom.get("avg_position", 0))
        ws4.cell(row=row, column=5, value=dom.get("unique_urls", 0))
        ws4.cell(row=row, column=6, value=", ".join(dom.get("keywords", [])))
        apply_row_format(ws4, row, 6)

    for i, width in enumerate([5, 35, 12, 18, 12, 80], 1):
        ws4.column_dimensions[get_column_letter(i)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"serp_analysis_{timestamp}.xlsx"

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


if __name__ == "__main__":
    print("\n" + "=" * 60)
    print("  SERP Analyzer - Versión Playwright v3")
    print("  Abre http://localhost:5000 en tu navegador")
    print("=" * 60 + "\n")
    app.run(debug=True, host="0.0.0.0", port=5000)
